"""Excel workbook for the early-warning layer. Counts and alerts are live formulas driven by the Settings sheet,
so the district team can change the date, window or thresholds in Excel and see the alerts update."""
import datetime as dt

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .early_warning import ISSUE_NAMES, LANG_NAMES

ARIAL = "Arial"
BLUE = Font(name=ARIAL, color="0000FF")
BOLD = Font(name=ARIAL, bold=True)
NORMAL = Font(name=ARIAL)
HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")


def _header(ws, cols, widths):
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        cell = ws.cell(row=1, column=i, value=c)
        cell.font, cell.fill = BOLD, HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def build(cfg, items, villages, findings, path):
    wb = Workbook()
    # ---------------------------------------------------------------- About
    ab = wb.active
    ab.title = "About"
    lines = [("Early-warning workbook: Use case 8, AIGP26 capstone", BOLD),
             ("All data is synthetic (made up). Village names are invented; any match with a real village is coincidental.", NORMAL),
             ("", NORMAL),
             ("How to use", BOLD),
             ("1. Change the yellow cells on the Settings sheet (date, window, thresholds). Blue text = values you can change.", NORMAL),
             ("2. 'Issue trends' recalculates counts, change and alerts for every issue and block.", NORMAL),
             ("3. 'Village counts' shows complaints per village in the last window; counts below the privacy limit show as <5.", NORMAL),
             ("4. 'Complaints' holds the 500 complaints with the system's issue and language for each.", NORMAL),
             ("5. 'Checks' lists the automated early-warning checks and their results (from the prototype run).", NORMAL),
             ("", NORMAL),
             ("An alert is a recommendation for field verification. It never triggers action by itself.", BOLD),
             ("Complaints still awaiting officer confirmation are counted, using the system's suggested issue, so that "
              "local-language complaints (which go to officers most often) are not under-counted.", NORMAL)]
    for i, (t, f) in enumerate(lines, 1):
        ab.cell(row=i, column=1, value=t).font = f
    ab.column_dimensions["A"].width = 120

    # ---------------------------------------------------------------- Settings
    st = wb.create_sheet("Settings")
    st["A1"], st["A1"].font = "Setting", BOLD
    st["B1"], st["B1"].font = "Value", BOLD
    st["C1"], st["C1"].font = "Meaning", BOLD
    rules, high = cfg.get("alert_rules", {}), cfg.get("high_priority", {})
    inputs = [("Alert date (as of)", dt.date.fromisoformat(str(cfg["as_of"])), "Day the alerts are produced", "yyyy-mm-dd"),
              ("Window (days)", int(cfg.get("window_days", 14)), "Last N days compared with the N days before", "0"),
              ("Alert: minimum complaints", int(rules.get("min_complaints", 20)), "Fewer than this never alerts", "0"),
              ("Alert: minimum rise", int(rules.get("min_rise_pct", 50)) / 100, "Rise vs previous window needed to alert", "0%"),
              ("High priority: minimum complaints", int(high.get("min_complaints", 50)), "Alert becomes High at or above this", "0"),
              ("High priority: minimum rise", int(high.get("min_rise_pct", 50)) / 100, "and at or above this rise", "0%"),
              ("Privacy: hide village counts below", int(cfg.get("suppress_below", 5)), "Shown as <N to protect individuals", "0")]
    for i, (label, val, meaning, fmt) in enumerate(inputs, 2):
        st.cell(row=i, column=1, value=label).font = NORMAL
        c = st.cell(row=i, column=2, value=val)
        c.font, c.fill, c.number_format = BLUE, INPUT_FILL, fmt
        st.cell(row=i, column=3, value=meaning).font = NORMAL
    st["B2"].comment = Comment("Source: config/early_warning.yaml (as_of). Synthetic demo date.", "AIGP26")
    derived = [("Last window: first day", "=B2-B3+1"), ("Last window: last day", "=B2"),
               ("Previous window: first day", "=B2-2*B3+1"), ("Previous window: last day", "=B2-B3")]
    for i, (label, f) in enumerate(derived, 10):
        st.cell(row=i, column=1, value=label).font = NORMAL
        c = st.cell(row=i, column=2, value=f)
        c.font, c.number_format = NORMAL, "yyyy-mm-dd"
    st.column_dimensions["A"].width, st.column_dimensions["B"].width, st.column_dimensions["C"].width = 36, 14, 48
    S = {"min": "Settings!$B$4", "rise": "Settings!$B$5", "hmin": "Settings!$B$6", "hrise": "Settings!$B$7",
         "sup": "Settings!$B$8", "c0": "Settings!$B$10", "c1": "Settings!$B$11", "p0": "Settings!$B$12", "p1": "Settings!$B$13"}

    # ---------------------------------------------------------------- Complaints
    cp = wb.create_sheet("Complaints")
    cols = ["Complaint id", "Date", "Block", "Village", "Channel", "Language (true)", "Language (detected)",
            "Issue (system)", "Awaiting officer", "Sensitive", "True issue (test label)", "Scenario (test label)", "Complaint text"]
    _header(cp, cols, [12, 11, 12, 14, 9, 14, 15, 13, 10, 9, 15, 17, 70])
    n = len(items)
    for r, it in enumerate(items, 2):
        vals = [it["complaint_id"], dt.date.fromisoformat(it["date"]), it["block"], it["village"], it["channel"],
                LANG_NAMES.get(it["language_group"], it["language_group"]),
                LANG_NAMES.get(it["detected_language"], it["detected_language"]), it["issue"],
                "yes" if it["provisional"] else "no", "yes" if it["sensitive"] else "no", it["true_issue"], it["scenario"], it["text"]]
        for c, v in enumerate(vals, 1):
            cell = cp.cell(row=r, column=c, value=v)
            cell.font = NORMAL
            if c == 2:
                cell.number_format = "yyyy-mm-dd"
    last = n + 1
    rng = lambda col: f"Complaints!${col}$2:${col}${last}"
    D, B, ISS, V, PROV = rng("B"), rng("C"), rng("H"), rng("D"), rng("I")

    # ---------------------------------------------------------------- Issue trends
    tr = wb.create_sheet("Issue trends")
    _header(tr, ["Issue code", "Issue", "Block", "Last window", "Previous window", "Change", "Awaiting officer (last window)",
                 "Alert", "Priority"], [11, 16, 13, 12, 14, 10, 16, 9, 10])
    blocks = sorted({v["block"] for v in villages})
    issues = [k for k in ISSUE_NAMES if k != "police"]
    r = 2
    for iss in issues:
        for blk in blocks:
            tr.cell(row=r, column=1, value=iss)
            tr.cell(row=r, column=2, value=ISSUE_NAMES[iss])
            tr.cell(row=r, column=3, value=blk)
            tr.cell(row=r, column=4, value=f'=COUNTIFS({ISS},A{r},{B},C{r},{D},">="&{S["c0"]},{D},"<="&{S["c1"]})')
            tr.cell(row=r, column=5, value=f'=COUNTIFS({ISS},A{r},{B},C{r},{D},">="&{S["p0"]},{D},"<="&{S["p1"]})')
            tr.cell(row=r, column=6, value=f'=IF(E{r}=0,"",(D{r}-E{r})/E{r})').number_format = "0%"
            tr.cell(row=r, column=7, value=f'=COUNTIFS({ISS},A{r},{B},C{r},{D},">="&{S["c0"]},{D},"<="&{S["c1"]},{PROV},"yes")')
            tr.cell(row=r, column=8, value=f'=IF(AND(D{r}>={S["min"]},OR(E{r}=0,F{r}>={S["rise"]})),"ALERT","")')
            tr.cell(row=r, column=9, value=f'=IF(H{r}="","",IF(AND(D{r}>={S["hmin"]},OR(E{r}=0,F{r}>={S["hrise"]})),"High","Medium"))')
            for c in range(1, 10):
                tr.cell(row=r, column=c).font = NORMAL
            r += 1
    tr.auto_filter.ref = f"A1:I{r - 1}"

    # ---------------------------------------------------------------- Village counts
    vc = wb.create_sheet("Village counts")
    _header(vc, ["Village", "Village (Devanagari)", "Block"] + [ISSUE_NAMES[i] for i in issues], [14, 18, 13] + [13] * len(issues))
    vc.cell(row=1, column=3 + len(issues) + 2, value="Counts are complaints in the last window. Below the privacy limit "
            "they show as <N; 0 shows as 0.").font = Font(name=ARIAL, italic=True)
    for r, v in enumerate(villages, 2):
        vc.cell(row=r, column=1, value=v["village"]).font = NORMAL
        vc.cell(row=r, column=2, value=v.get("village_local", "")).font = NORMAL
        vc.cell(row=r, column=3, value=v["block"]).font = NORMAL
        for j, iss in enumerate(issues, 4):
            cnt = f'COUNTIFS({ISS},"{iss}",{V},$A{r},{D},">="&{S["c0"]},{D},"<="&{S["c1"]})'
            vc.cell(row=r, column=j, value=f'=IF({cnt}=0,0,IF({cnt}<{S["sup"]},"<"&{S["sup"]},{cnt}))').font = NORMAL

    # ---------------------------------------------------------------- Villages
    vl = wb.create_sheet("Villages")
    _header(vl, ["Village", "Village (Devanagari)", "Block", "District", "Settlement"], [14, 18, 13, 15, 12])
    for r, v in enumerate(villages, 2):
        for c, k in enumerate(["village", "village_local", "block", "district", "settlement"], 1):
            vl.cell(row=r, column=c, value=v.get(k, "")).font = NORMAL

    # ---------------------------------------------------------------- Checks
    ckw = wb.create_sheet("Checks")
    _header(ckw, ["Result", "Severity if failed", "Check", "What was tested", "Expected", "What happened"], [8, 12, 40, 40, 34, 70])
    for r, f in enumerate(sorted(findings, key=lambda f: f.passed), 2):
        for c, v in enumerate(["Pass" if f.passed else "Fail", f.severity, f.probe, f.tested, f.expected, f.actual], 1):
            cell = ckw.cell(row=r, column=c, value=v)
            cell.font, cell.alignment = NORMAL, Alignment(wrap_text=True, vertical="top")
    ckw.cell(row=len(findings) + 3, column=1,
             value="Results of the automated early-warning checks from the prototype run (python -m grievdesk early-warning). "
                   "Fixed values, not formulas.").font = Font(name=ARIAL, italic=True)
    wb.save(path)
    return path
