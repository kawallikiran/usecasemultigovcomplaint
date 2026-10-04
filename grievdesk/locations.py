"""Real places: zone, state, district and city/town, from the Census 2011 list of cities and towns
(data/grievance/source_census2011_towns.xlsx), brought to current state boundaries.

Changes applied to the 2011 list:
  * names: Orissa -> Odisha, Uttaranchal -> Uttarakhand, Pondicherry -> Puducherry
  * Telangana (formed 2014): its ten 2011 districts move from Andhra Pradesh to Telangana
  * Ladakh (formed 2019): Leh (Ladakh) and Kargil move from Jammu & Kashmir to Ladakh
  * Dadra & Nagar Haveli and Daman & Diu (merged 2020) become one union territory
  * duplicate rows and the notes at the end of the sheet are removed
District names stay as in 2011; some districts have since been split or renamed.
"""
import csv
import os

DATA = os.path.join(os.path.dirname(__file__), os.pardir, "data", "grievance")
SOURCE = os.path.join(DATA, "source_census2011_towns.xlsx")
LOCATIONS = os.path.join(DATA, "locations.csv")

ZONES = {
    "North Zone": ["Chandigarh", "Delhi", "Haryana", "Himachal Pradesh", "Jammu & Kashmir", "Ladakh", "Punjab", "Rajasthan"],
    "South Zone": ["Andaman & Nicobar", "Andhra Pradesh", "Karnataka", "Kerala", "Lakshadweep", "Puducherry", "Tamil Nadu",
                   "Telangana"],
    "East Zone": ["Bihar", "Jharkhand", "Odisha", "West Bengal"],
    "West Zone": ["Dadra and Nagar Haveli and Daman and Diu", "Goa", "Gujarat", "Maharashtra"],
    "Central Zone": ["Chhattisgarh", "Madhya Pradesh", "Uttar Pradesh", "Uttarakhand"],
    "Northeast Zone": ["Arunachal Pradesh", "Assam", "Manipur", "Meghalaya", "Mizoram", "Nagaland", "Sikkim", "Tripura"],
}
ZONE_OF = {s: z for z, states in ZONES.items() for s in states}
RENAME = {"Orissa": "Odisha", "Uttaranchal": "Uttarakhand", "Pondicherry": "Puducherry",
          "Andaman & Nicobar Islands": "Andaman & Nicobar", "Dadra & Nagar Haveli": "Dadra and Nagar Haveli and Daman and Diu",
          "Daman & Diu": "Dadra and Nagar Haveli and Daman and Diu"}
TELANGANA_2011 = {"Adilabad", "Nizamabad", "Karimnagar", "Medak", "Hyderabad", "Rangareddy", "Mahbubnagar", "Nalgonda",
                  "Warangal", "Khammam", "Rangareddi", "Hyderabad and Rangareddi"}
LADAKH_2011 = {"Leh (Ladakh)", "Leh(Ladakh)", "Leh", "Kargil"}
STATUS = {"C.T.": "Census Town", "M": "Municipality", "N.P.": "Nagar Panchayat", "T.P.": "Town Panchayat",
          "M.Cl.": "Municipal Council", "M.B.": "Municipal Board", "N.A.C.": "Notified Area Committee",
          "M.Corp.": "Municipal Corporation", "N.A.": "Notified Area", "T.M.C": "Town Municipal Council",
          "T.C.": "Town Committee", "M.Crop.": "Municipal Corporation", "M.C.": "Municipal Committee", "C.M.C.": "City Municipal Council",
          "C.B.": "Cantonment Board", "N.T.": "Notified Town", "I.N.A.": "Industrial Notified Area"}


def build():
    """Reads the Census sheet and writes data/grievance/locations.csv. Returns the number of places."""
    from openpyxl import load_workbook
    ws = load_workbook(SOURCE, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    code = lambda c: str(c or "").replace(".", "").replace(" ", "").upper()
    legend = {code(r[1]): str(r[3]).split("/")[0].strip().rstrip(".") for r in rows
              if r and r[1] and not r[4] and r[2] == ":" and r[3]}            # the sheet's own key to status codes
    seen, out = set(), []
    for r in rows:
        if not r or not r[1] or not r[4] or not r[6]:
            continue
        town, status, state, district = str(r[1]).strip(), str(r[2] or "").strip(), str(r[4]).replace("*", "").strip(), str(r[6]).strip()
        state = RENAME.get(state, state)
        if state == "Andhra Pradesh" and district in TELANGANA_2011:
            state = "Telangana"
        if state == "Jammu & Kashmir" and district in LADAKH_2011:
            state = "Ladakh"
        key = (state, district, town)
        if key in seen:
            continue
        seen.add(key)
        out.append({"zone": ZONE_OF.get(state, "Unmapped"), "state": state, "district": district, "town": town,
                    "urban_status": STATUS.get(status) or legend.get(code(status), status), "census_state_code": r[3], "census_district_code": r[5]})
    out.sort(key=lambda x: (x["zone"], x["state"], x["district"], x["town"]))
    with open(LOCATIONS, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(out[0]))
        w.writeheader()
        w.writerows(out)
    return len(out)


def load():
    with open(LOCATIONS, encoding="utf-8") as f:
        return list(csv.DictReader(f))


def tree():
    """{state: {district: [town, ...]}} for dropdowns (states and districts sorted)."""
    t = {}
    for r in load():
        t.setdefault(r["state"], {}).setdefault(r["district"], []).append(r["town"])
    return {s: dict(sorted(d.items())) for s, d in sorted(t.items())}


def zone_of(state):
    return ZONE_OF.get(state, "Unmapped")
